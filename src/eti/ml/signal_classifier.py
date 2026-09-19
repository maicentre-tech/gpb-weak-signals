"""Воспроизводимый baseline-классификатор слабых технологических сигналов.

Модель намеренно не использует LLM. Она объединяет текстовые поля, явно
заданные методологами, и короткие структурные маркеры стадии и динамики.
Каждое решение можно разложить на вклад конкретных токенов, поэтому этот
контур пригоден и для оценки на скрытой выборке, и для демонстрации жюри.
"""

from __future__ import annotations

from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

import joblib
from openpyxl import load_workbook
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import accuracy_score, classification_report, precision_recall_fscore_support
from sklearn.model_selection import train_test_split


@dataclass(frozen=True, slots=True)
class SignalRecord:
    technology: str
    domain: str
    rationale: str
    stage: str
    mention_trend: str
    sources: str
    label: int

    def text(self) -> str:
        """Формат, который одновременно читает vectorizer и аналитик."""
        return " ".join(
            (
                self.technology,
                self.rationale,
                f"domain_{self.domain.lower().replace(' ', '_')}",
                f"stage_{stage_bucket(self.stage)}",
                f"trend_{trend_bucket(self.mention_trend)}",
                f"sources_{source_bucket(self.sources)}",
            )
        )


def stage_bucket(value: str) -> str:
    value = value.lower()
    if "массов" in value or "зрел" in value or "стандарт" in value:
        return "mature"
    if "раннее внедрение" in value:
        return "early_adoption"
    if "пилот" in value:
        return "pilot"
    if "прототип" in value or "poc" in value:
        return "prototype"
    return "research"


def trend_bucket(value: str) -> str:
    value = value.lower()
    if "растёт быстро" in value or "ускор" in value:
        return "accelerating"
    if "растёт" in value:
        return "growing"
    if "стабиль" in value:
        return "flat"
    if "падает" in value or "снижа" in value:
        return "declining"
    return "unknown"


def source_bucket(value: str) -> str:
    count = value.count("http")
    return "multi" if count >= 2 else "single" if count == 1 else "none"


def load_positive_signals(path: Path) -> list[SignalRecord]:
    """Читает предоставленный xlsx, не полагаясь на названия столбцов.

    В файле первая строка — заголовок листа, вторая — реальные заголовки.
    Все записи из него являются положительными экспертными примерами.
    """
    sheet = load_workbook(path, data_only=True, read_only=True).active
    rows = list(sheet.iter_rows(min_row=3, values_only=True))
    records: list[SignalRecord] = []
    for row in rows:
        if not isinstance(row[1], int):
            continue
        records.append(
            SignalRecord(
                technology=str(row[2] or ""), domain=str(row[3] or ""),
                rationale=str(row[5] or ""), stage=str(row[6] or ""),
                mention_trend=str(row[7] or ""), sources=str(row[9] or ""), label=1,
            )
        )
    if len(records) != 100:
        raise ValueError(f"Ожидалось 100 положительных сигналов, получено {len(records)}")
    return records


# Негативы не маскируются под разметку организаторов: это открытый и
# редактируемый набор контрпримеров для трёх обязательных классов отсечения.
NEGATIVE_SEEDS: tuple[SignalRecord, ...] = tuple(
    SignalRecord(*row, label=0)
    for row in (
        ("Kubernetes для корпоративной инфраструктуры", "Инфраструктура ИИ", "Сформированный рынок, широкое внедрение и устойчивые лидеры.", "Массовое внедрение", "Стабильный высокий объём упоминаний", "https://kubernetes.io"),
        ("REST API", "Индустриальный ИИ", "Отраслевой стандарт с многолетней практикой и зрелой экосистемой.", "Зрелая технология", "Стабильный", "https://www.ietf.org"),
        ("Контейнеризация Docker", "Инфраструктура ИИ", "Массово используемая технология с устойчивой конкурентной структурой.", "Массовое внедрение", "Стабильный", "https://www.docker.com"),
        ("Облачное объектное хранилище", "Инфраструктура ИИ", "Сформированный рынок с гиперскейлерами и типовыми закупками.", "Зрелая технология", "Стабильный", "https://aws.amazon.com/s3"),
        ("Традиционный VPN", "Защита ИИ", "Зрелый массовый стандарт доступа, не ранний индикатор.", "Зрелая технология", "Стабильный", "https://www.nist.gov"),
        ("Реляционная СУБД", "Финтех", "Технология имеет зрелый рынок, стандарты и массовое внедрение.", "Массовое внедрение", "Стабильный", "https://www.postgresql.org"),
        ("Блокчейн для всего", "Финтех", "Маркетинговый анонс без независимой научной, патентной или внедренческой верификации.", "Концепция", "Растёт быстро в пресс-релизах", "https://example.org/press"),
        ("Революционный ИИ-стартап без продукта", "Индустриальный ИИ", "Один рекламный материал, нет независимых источников и технического описания.", "Концепция", "Растёт быстро", "https://example.org/press"),
        ("Квантовый маркетинговый токен", "Финтех", "Высокая медийность при отсутствии публикаций, патентов и воспроизводимых кейсов.", "Концепция", "Растёт быстро", "https://example.org/press"),
        ("Универсальный AGI-плагин", "Индустриальный ИИ", "Вендорское обещание без подтверждённого поиска и независимых источников.", "Концепция", "Растёт быстро", "https://example.org/press"),
        ("Инфлюенсерский робот-помощник", "Роботы", "Единичное вирусное видео, нет технической документации или независимой проверки.", "Концепция", "Растёт быстро", "https://example.org/video"),
        ("Нейросеть для чтения мыслей по селфи", "Защита ИИ", "Рекламная формулировка не подкреплена исследованием или патентом.", "Концепция", "Растёт быстро", "https://example.org/post"),
        ("Небольшой GitHub-проект без релизов", "Edge", "Один репозиторий с малой активностью и без кросс-источникового подтверждения.", "Прототип", "Неизвестно", "https://github.com/example/project"),
        ("Одиночный препринт без продолжения", "Роботы", "Единичная публикация, без репликаций, патентов или внедренческих следов.", "Исследование", "Стабильный", "https://arxiv.org"),
        ("Локальный хакатонный бот", "Индустриальный ИИ", "Малый объём, один источник и отсутствует динамика.", "Прототип", "Неизвестно", "https://example.org/demo"),
        ("Нишевый форк криптобиблиотеки", "Финтех", "Малый технический артефакт без независимых упоминаний и роста.", "Прототип", "Стабильный", "https://github.com/example/fork"),
        ("Единичный патент без семейства", "Edge", "Одна заявка без роста, цитирования или подтверждения из иных источников.", "Исследование", "Стабильный", "https://patents.google.com"),
        ("Тестовый датчик в лаборатории", "Роботы", "Одна лабораторная новость без повторных публикаций и внедрения.", "Прототип", "Неизвестно", "https://example.edu/news"),
    )
)


class SignalClassifier:
    """TF-IDF + логистическая регрессия с локальным объяснением решения."""

    def __init__(self) -> None:
        self.vectorizer = TfidfVectorizer(ngram_range=(1, 2), min_df=1, sublinear_tf=True)
        self.model = LogisticRegression(max_iter=2_000, class_weight="balanced", random_state=42)

    def fit(self, records: list[SignalRecord]) -> SignalClassifier:
        self.vectorizer.fit(record.text() for record in records)
        matrix = self.vectorizer.transform(record.text() for record in records)
        self.model.fit(matrix, [record.label for record in records])
        return self

    def predict_with_explanation(self, record: SignalRecord, top_k: int = 8) -> dict[str, Any]:
        matrix = self.vectorizer.transform([record.text()])
        probability = float(self.model.predict_proba(matrix)[0, 1])
        names = self.vectorizer.get_feature_names_out()
        contributions = matrix.toarray()[0] * self.model.coef_[0]
        ranked = sorted(range(len(contributions)), key=lambda i: abs(contributions[i]), reverse=True)
        predictors = [
            {"feature": str(names[i]), "contribution": round(float(contributions[i]), 4)}
            for i in ranked[:top_k]
            if contributions[i] != 0
        ]
        return {
            "weak_signal_probability": round(probability, 4),
            "label": "weak_signal" if probability >= 0.5 else "not_weak_signal",
            "key_predictors": predictors,
        }

    def evaluate(self, records: list[SignalRecord]) -> dict[str, Any]:
        labels = [record.label for record in records]
        train, test = train_test_split(records, test_size=0.25, stratify=labels, random_state=42)
        self.fit(train)
        predicted = self.model.predict(self.vectorizer.transform(record.text() for record in test))
        precision, recall, f1, _ = precision_recall_fscore_support(
            [r.label for r in test], predicted, average="binary", zero_division=0
        )
        return {
            "split": {"random_state": 42, "train": len(train), "test": len(test)},
            "accuracy": round(float(accuracy_score([r.label for r in test], predicted)), 4),
            "precision": round(float(precision), 4),
            "recall": round(float(recall), 4),
            "f1": round(float(f1), 4),
            "report": classification_report([r.label for r in test], predicted, output_dict=True, zero_division=0),
            "warning": "Метрики относятся к открытому baseline-набору: 100 экспертных положительных примеров и явно маркированные контрпримеры. Финальный порог принимается только по закрытой выборке организаторов.",
        }

    def save(self, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        joblib.dump({"vectorizer": self.vectorizer, "model": self.model}, path)

    @classmethod
    def load(cls, path: Path) -> SignalClassifier:
        saved = joblib.load(path)
        instance = cls()
        instance.vectorizer, instance.model = saved["vectorizer"], saved["model"]
        return instance


def records_as_json(records: list[SignalRecord]) -> list[dict[str, Any]]:
    return [asdict(record) for record in records]
