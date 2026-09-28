"""Воспроизводимый baseline-классификатор слабых технологических сигналов.

Модель намеренно не использует LLM. Она объединяет текстовые поля, явно
заданные методологами, и короткие структурные маркеры стадии и динамики.
Каждое решение можно разложить на вклад конкретных токенов, поэтому этот
контур пригоден и для оценки на скрытой выборке, и для демонстрации жюри.
"""

from __future__ import annotations

from collections import Counter
from dataclasses import asdict, dataclass, field
from pathlib import Path
import re
from typing import Any
import unicodedata

import joblib
import numpy as np
from openpyxl import load_workbook
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import (
    accuracy_score,
    average_precision_score,
    balanced_accuracy_score,
    brier_score_loss,
    classification_report,
    confusion_matrix,
    f1_score,
    precision_score,
    recall_score,
)
from sklearn.model_selection import StratifiedGroupKFold
from sklearn.pipeline import Pipeline


DEFAULT_FEATURE_GROUPS = frozenset({"technology", "maturity", "trend"})
ALL_FEATURE_GROUPS = frozenset(
    {"technology", "rationale", "domain", "maturity", "trend", "sources"}
)
MIN_KEY_PREDICTORS = 3
MAX_KEY_PREDICTORS = 6
EVALUATION_VERSION = "eti-grouped-nested-cv-v1"


def normalize_technology_name(value: str) -> str:
    """Нормализует имя для проверки дублей и группировки folds."""
    normalized = unicodedata.normalize("NFKC", value).casefold()
    return " ".join(re.findall(r"[\w]+", normalized, flags=re.UNICODE))


@dataclass(frozen=True, slots=True)
class SignalRecord:
    technology: str
    domain: str
    rationale: str
    stage: str
    mention_trend: str
    sources: str
    label: int
    example_id: str = ""
    group_id: str = ""
    label_source: str = "unspecified"
    provenance: dict[str, Any] = field(default_factory=dict)

    def text(self, feature_groups: frozenset[str] | set[str] | None = None) -> str:
        """Возвращает только выбранные группы модели; provenance не становится признаком.

        По умолчанию используются признаки, доступные и в production scoring.
        Свободный rationale, domain и source count остаются для аудита, но не
        обучают baseline:
        в production rationale содержит числовые score, а domain сейчас является
        служебной константой. Число evidence-ссылок пока несопоставимо с
        production placeholder.
        """
        groups = DEFAULT_FEATURE_GROUPS if feature_groups is None else feature_groups
        unknown = set(groups) - ALL_FEATURE_GROUPS
        if unknown:
            raise ValueError(f"Неизвестные группы признаков: {sorted(unknown)}")

        parts: list[str] = []
        if "technology" in groups:
            parts.append(self.technology)
        if "rationale" in groups:
            parts.append(self.rationale)
        if "domain" in groups:
            parts.append(f"domain_{self.domain.lower().replace(' ', '_')}")
        if "maturity" in groups:
            parts.append(f"stage_{stage_bucket(self.stage)}")
        if "trend" in groups:
            parts.append(f"trend_{trend_bucket(self.mention_trend)}")
        if "sources" in groups:
            parts.append(f"sources_{source_bucket(self.sources)}")
        return " ".join(part for part in parts if part)


def stage_bucket(value: str) -> str:
    value = value.lower()
    if "массов" in value or "зрел" in value or "стандарт" in value:
        return "mature"
    if "раннее внедрение" in value:
        return "early_adoption"
    if "пилот" in value:
        return "pilot"
    if "прототип" in value or "poc" in value:
        return "prototype"
    return "research"


def trend_bucket(value: str) -> str:
    value = value.lower()
    if "растёт быстро" in value or "ускор" in value:
        return "accelerating"
    if "растёт" in value:
        return "growing"
    if "стабиль" in value:
        return "flat"
    if "падает" in value or "снижа" in value:
        return "declining"
    return "unknown"


def source_bucket(value: str) -> str:
    count = value.count("http")
    return "multi" if count >= 2 else "single" if count == 1 else "none"


def load_positive_signals(path: Path) -> list[SignalRecord]:
    """Читает предоставленный xlsx, не полагаясь на названия столбцов.

    В файле первая строка — заголовок листа, вторая — реальные заголовки.
    Все записи из него являются положительными экспертными примерами.
    """
    sheet = load_workbook(path, data_only=True, read_only=True).active
    rows = list(sheet.iter_rows(min_row=3, values_only=True))
    records: list[SignalRecord] = []
    for row_number, row in enumerate(rows, start=3):
        if not isinstance(row[1], int):
            continue
        technology = str(row[2] or "")
        records.append(
            SignalRecord(
                technology=technology, domain=str(row[3] or ""),
                rationale=str(row[5] or ""), stage=str(row[6] or ""),
                mention_trend=str(row[7] or ""), sources=str(row[9] or ""), label=1,
                example_id=f"benchmark-xlsx-row-{row_number}",
                group_id=normalize_technology_name(technology),
                label_source="expert_benchmark",
                provenance={"source_file": path.name, "row": row_number},
            )
        )
    if len(records) != 100:
        raise ValueError(f"Ожидалось 100 положительных сигналов, получено {len(records)}")
    return records


# Негативы не маскируются под разметку организаторов: это открытый и
# редактируемый набор контрпримеров для трёх обязательных классов отсечения.
_NEGATIVE_SEED_ROWS = (
        ("Kubernetes для корпоративной инфраструктуры", "Инфраструктура ИИ", "Сформированный рынок, широкое внедрение и устойчивые лидеры.", "Массовое внедрение", "Стабильный высокий объём упоминаний", "https://kubernetes.io"),
        ("REST API", "Индустриальный ИИ", "Отраслевой стандарт с многолетней практикой и зрелой экосистемой.", "Зрелая технология", "Стабильный", "https://www.ietf.org"),
        ("Контейнеризация Docker", "Инфраструктура ИИ", "Массово используемая технология с устойчивой конкурентной структурой.", "Массовое внедрение", "Стабильный", "https://www.docker.com"),
        ("Облачное объектное хранилище", "Инфраструктура ИИ", "Сформированный рынок с гиперскейлерами и типовыми закупками.", "Зрелая технология", "Стабильный", "https://aws.amazon.com/s3"),
        ("Традиционный VPN", "Защита ИИ", "Зрелый массовый стандарт доступа, не ранний индикатор.", "Зрелая технология", "Стабильный", "https://www.nist.gov"),
        ("Реляционная СУБД", "Финтех", "Технология имеет зрелый рынок, стандарты и массовое внедрение.", "Массовое внедрение", "Стабильный", "https://www.postgresql.org"),
        ("Блокчейн для всего", "Финтех", "Маркетинговый анонс без независимой научной, патентной или внедренческой верификации.", "Концепция", "Растёт быстро в пресс-релизах", "https://example.org/press"),
        ("Революционный ИИ-стартап без продукта", "Индустриальный ИИ", "Один рекламный материал, нет независимых источников и технического описания.", "Концепция", "Растёт быстро", "https://example.org/press"),
        ("Квантовый маркетинговый токен", "Финтех", "Высокая медийность при отсутствии публикаций, патентов и воспроизводимых кейсов.", "Концепция", "Растёт быстро", "https://example.org/press"),
        ("Универсальный AGI-плагин", "Индустриальный ИИ", "Вендорское обещание без подтверждённого поиска и независимых источников.", "Концепция", "Растёт быстро", "https://example.org/press"),
        ("Инфлюенсерский робот-помощник", "Роботы", "Единичное вирусное видео, нет технической документации или независимой проверки.", "Концепция", "Растёт быстро", "https://example.org/video"),
        ("Нейросеть для чтения мыслей по селфи", "Защита ИИ", "Рекламная формулировка не подкреплена исследованием или патентом.", "Концепция", "Растёт быстро", "https://example.org/post"),
        ("Небольшой GitHub-проект без релизов", "Edge", "Один репозиторий с малой активностью и без кросс-источникового подтверждения.", "Прототип", "Неизвестно", "https://github.com/example/project"),
        ("Одиночный препринт без продолжения", "Роботы", "Единичная публикация, без репликаций, патентов или внедренческих следов.", "Исследование", "Стабильный", "https://arxiv.org"),
        ("Локальный хакатонный бот", "Индустриальный ИИ", "Малый объём, один источник и отсутствует динамика.", "Прототип", "Неизвестно", "https://example.org/demo"),
        ("Нишевый форк криптобиблиотеки", "Финтех", "Малый технический артефакт без независимых упоминаний и роста.", "Прототип", "Стабильный", "https://github.com/example/fork"),
        ("Единичный патент без семейства", "Edge", "Одна заявка без роста, цитирования или подтверждения из иных источников.", "Исследование", "Стабильный", "https://patents.google.com"),
        ("Тестовый датчик в лаборатории", "Роботы", "Одна лабораторная новость без повторных публикаций и внедрения.", "Прототип", "Неизвестно", "https://example.edu/news"),
)
NEGATIVE_SEEDS: tuple[SignalRecord, ...] = tuple(
    SignalRecord(
        *row,
        label=0,
        example_id=f"illustrative-seed-{index:02d}",
        group_id=normalize_technology_name(row[0]),
        label_source="illustrative_seed",
        provenance={"review_status": "curated_example"},
    )
    for index, row in enumerate(_NEGATIVE_SEED_ROWS, start=1)
)


class SignalClassifier:
    """TF-IDF + логистическая регрессия с локальным объяснением решения."""

    def __init__(self) -> None:
        self.vectorizer = TfidfVectorizer(ngram_range=(1, 2), min_df=1, sublinear_tf=True)
        self.model = LogisticRegression(max_iter=2_000, class_weight="balanced", random_state=42)
        self.calibrator: LogisticRegression | None = None
        self.decision_threshold = 0.5
        self.calibration_method = "not_calibrated"
        self.feature_groups = DEFAULT_FEATURE_GROUPS

    def fit(self, records: list[SignalRecord]) -> SignalClassifier:
        _validate_training_records(records, require_group_cv=False)
        texts = [record.text(self.feature_groups) for record in records]
        self.vectorizer.fit(texts)
        matrix = self.vectorizer.transform(texts)
        self.model.fit(matrix, [record.label for record in records])
        self.calibrator = None
        self.calibration_method = "not_calibrated"
        return self

    def fit_calibrated(
        self, records: list[SignalRecord], *, n_splits: int = 5, random_state: int = 42
    ) -> SignalClassifier:
        """Fit final model and a sigmoid calibrator on grouped out-of-fold scores."""
        _validate_training_records(records)
        labels = np.asarray([record.label for record in records], dtype=int)
        raw_oof, _, actual_splits = _oof_raw_probabilities(
            records, self.feature_groups, n_splits=n_splits, random_state=random_state
        )
        self.fit(records)
        self.calibrator = _fit_sigmoid_calibrator(raw_oof, labels)
        self.calibration_method = f"sigmoid_grouped_oof_{actual_splits}_fold"
        return self

    def predict_with_explanation(
        self, record: SignalRecord, top_k: int = MAX_KEY_PREDICTORS
    ) -> dict[str, Any]:
        top_k = min(MAX_KEY_PREDICTORS, max(MIN_KEY_PREDICTORS, int(top_k)))
        matrix = self.vectorizer.transform([record.text(self.feature_groups)])
        raw_probability = float(self.model.predict_proba(matrix)[0, 1])
        probability = (
            _calibrated_probabilities(self.calibrator, np.asarray([raw_probability]))[0]
            if self.calibrator is not None
            else raw_probability
        )
        names = self.vectorizer.get_feature_names_out()
        contributions = matrix.toarray()[0] * self.model.coef_[0]
        ranked = sorted(range(len(contributions)), key=lambda i: abs(contributions[i]), reverse=True)
        predictors = [
            {"feature": str(names[i]), "contribution": round(float(contributions[i]), 4)}
            for i in ranked[:top_k]
            if contributions[i] != 0
        ]
        return {
            "weak_signal_probability": round(probability, 4),
            "raw_probability": round(raw_probability, 4),
            "calibrated": self.calibrator is not None,
            "calibration_method": self.calibration_method,
            "decision_threshold": self.decision_threshold,
            "predictor_basis": "raw linear-model contribution; calibration transforms probability only",
            "label": "weak_signal" if probability >= self.decision_threshold else "not_weak_signal",
            "key_predictors": predictors,
        }

    def evaluate(
        self, records: list[SignalRecord], *, n_splits: int = 5, random_state: int = 42
    ) -> dict[str, Any]:
        """Nested, grouped CV; calibration is learned only from each outer train fold."""
        _validate_training_records(records)
        primary = _nested_grouped_evaluation(
            records, DEFAULT_FEATURE_GROUPS, n_splits=n_splits, random_state=random_state
        )
        variants = {
            "without_technology_name": DEFAULT_FEATURE_GROUPS - {"technology"},
            "with_source_count": DEFAULT_FEATURE_GROUPS | {"sources"},
            "sources_only": frozenset({"sources"}),
            "domain_only": frozenset({"domain"}),
            "rationale_only": frozenset({"rationale"}),
        }
        ablations = {
            name: _nested_grouped_evaluation(
                records, groups, n_splits=n_splits, random_state=random_state
            )
            for name, groups in variants.items()
        }
        primary_ap = primary["calibrated"]["overall"]["average_precision"]
        shortcut_warnings = []
        for name in ("sources_only", "domain_only", "rationale_only"):
            audit_ap = ablations[name]["calibrated"]["overall"]["average_precision"]
            if audit_ap is not None and primary_ap is not None and audit_ap >= primary_ap - 0.05:
                shortcut_warnings.append(
                    f"{name} почти не уступает полной модели; проверьте возможную утечку "
                    "источника или способа разметки."
                )

        return {
            "evaluation_version": EVALUATION_VERSION,
            "protocol": "nested_stratified_group_cross_validation",
            "grouping": "stable group_id, otherwise normalized technology name",
            "outer_folds": primary["outer_folds"],
            "decision_threshold": self.decision_threshold,
            "calibration": "sigmoid fitted on inner out-of-fold scores per outer fold",
            "feature_groups": sorted(DEFAULT_FEATURE_GROUPS),
            "primary": primary,
            "feature_ablation": ablations,
            "shortcut_warnings": shortcut_warnings,
            "dataset_audit": audit_dataset(records),
            "warning": (
                "Это открытая кросс-валидация, а не независимая закрытая проверка. "
                "Корпусные кандидаты со статусом pending не являются подтверждённой "
                "разметкой и не должны попадать в обучение."
            ),
        }

    def save(self, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        joblib.dump(
            {
                "vectorizer": self.vectorizer,
                "model": self.model,
                "calibrator": self.calibrator,
                "decision_threshold": self.decision_threshold,
                "calibration_method": self.calibration_method,
                "feature_groups": sorted(self.feature_groups),
                "format_version": 2,
            },
            path,
        )

    @classmethod
    def load(cls, path: Path) -> SignalClassifier:
        saved = joblib.load(path)
        instance = cls()
        instance.vectorizer, instance.model = saved["vectorizer"], saved["model"]
        instance.calibrator = saved.get("calibrator")
        instance.decision_threshold = float(saved.get("decision_threshold", 0.5))
        instance.calibration_method = saved.get("calibration_method", "legacy_uncalibrated")
        instance.feature_groups = frozenset(saved.get("feature_groups", ALL_FEATURE_GROUPS))
        return instance


def records_as_json(records: list[SignalRecord]) -> list[dict[str, Any]]:
    return [asdict(record) for record in records]


def _validate_training_records(
    records: list[SignalRecord], *, require_group_cv: bool = True
) -> None:
    if len(records) < 2:
        raise ValueError("Для обучения нужны примеры обоих классов")
    labels = {record.label for record in records}
    if labels != {0, 1}:
        raise ValueError("Обучающая выборка должна содержать метки 0 и 1")

    group_labels: dict[str, set[int]] = {}
    name_labels: dict[str, set[int]] = {}
    for record in records:
        if record.label not in (0, 1):
            raise ValueError(f"Недопустимая метка: {record.label}")
        name = normalize_technology_name(record.technology)
        group = record.group_id.strip() or name
        if not group or not name:
            raise ValueError("У каждого примера должны быть technology и group_id")
        group_labels.setdefault(group, set()).add(record.label)
        name_labels.setdefault(name, set()).add(record.label)
    conflicts = [key for key, values in group_labels.items() if len(values) > 1]
    name_conflicts = [key for key, values in name_labels.items() if len(values) > 1]
    if conflicts or name_conflicts:
        raise ValueError(
            "Одна технология встречается с разными метками; разрешите конфликт "
            f"до обучения (group_id={conflicts[:3]}, technology={name_conflicts[:3]})"
        )

    counts = {
        label: len({record.group_id.strip() or normalize_technology_name(record.technology)
                    for record in records if record.label == label})
        for label in (0, 1)
    }
    if require_group_cv and min(counts.values()) < 2:
        raise ValueError("Для grouped CV нужны минимум две независимые группы каждого класса")


def _make_pipeline() -> Pipeline:
    return Pipeline(
        [
            ("tfidf", TfidfVectorizer(ngram_range=(1, 2), min_df=1, sublinear_tf=True)),
            ("model", LogisticRegression(max_iter=2_000, class_weight="balanced", random_state=42)),
        ]
    )


def _group_ids(records: list[SignalRecord]) -> np.ndarray:
    return np.asarray(
        [record.group_id.strip() or normalize_technology_name(record.technology) for record in records],
        dtype=object,
    )


def _fold_count(records: list[SignalRecord], requested: int) -> int:
    labels = np.asarray([record.label for record in records], dtype=int)
    groups = _group_ids(records)
    group_counts = [len(set(groups[labels == label])) for label in (0, 1)]
    actual = min(requested, *group_counts)
    if actual < 2:
        raise ValueError("Недостаточно независимых групп для grouped cross-validation")
    return actual


def _oof_raw_probabilities(
    records: list[SignalRecord],
    feature_groups: frozenset[str] | set[str],
    *,
    n_splits: int,
    random_state: int,
) -> tuple[np.ndarray, np.ndarray, int]:
    actual_splits = _fold_count(records, n_splits)
    labels = np.asarray([record.label for record in records], dtype=int)
    groups = _group_ids(records)
    texts = [record.text(feature_groups) for record in records]
    splitter = StratifiedGroupKFold(
        n_splits=actual_splits, shuffle=True, random_state=random_state
    )
    probabilities = np.full(len(records), np.nan, dtype=float)
    fold_ids = np.full(len(records), -1, dtype=int)
    for fold, (train_indices, test_indices) in enumerate(
        splitter.split(texts, labels, groups)
    ):
        if len(set(labels[train_indices])) < 2:
            raise ValueError(f"Fold {fold + 1} содержит в обучении один класс")
        estimator = _make_pipeline()
        estimator.fit([texts[i] for i in train_indices], labels[train_indices])
        probabilities[test_indices] = estimator.predict_proba(
            [texts[i] for i in test_indices]
        )[:, 1]
        fold_ids[test_indices] = fold
    if np.isnan(probabilities).any() or (fold_ids < 0).any():
        raise RuntimeError("Не для каждого примера получено out-of-fold предсказание")
    return probabilities, fold_ids, actual_splits


def _logit(values: np.ndarray) -> np.ndarray:
    clipped = np.clip(values, 1e-6, 1 - 1e-6)
    return np.log(clipped / (1 - clipped)).reshape(-1, 1)


def _fit_sigmoid_calibrator(
    probabilities: np.ndarray, labels: np.ndarray
) -> LogisticRegression:
    calibrator = LogisticRegression(max_iter=2_000, random_state=42)
    calibrator.fit(_logit(probabilities), labels)
    return calibrator


def _calibrated_probabilities(
    calibrator: LogisticRegression, probabilities: np.ndarray
) -> np.ndarray:
    return calibrator.predict_proba(_logit(probabilities))[:, 1]


def _calibration_error(labels: np.ndarray, probabilities: np.ndarray, bins: int = 10) -> float:
    bin_count = min(bins, max(2, len(labels)))
    bin_ids = np.minimum((np.clip(probabilities, 0, 1) * bin_count).astype(int), bin_count - 1)
    error = 0.0
    for bin_id in range(bin_count):
        in_bin = bin_ids == bin_id
        if in_bin.any():
            error += float(in_bin.mean()) * abs(
                float(labels[in_bin].mean()) - float(probabilities[in_bin].mean())
            )
    return error


def _metrics(labels: np.ndarray, probabilities: np.ndarray, threshold: float) -> dict[str, Any]:
    predicted = (probabilities >= threshold).astype(int)
    report = classification_report(
        labels, predicted, labels=[0, 1], output_dict=True, zero_division=0
    )
    return {
        "examples": int(len(labels)),
        "class_support": {"negative": int((labels == 0).sum()), "positive": int((labels == 1).sum())},
        "threshold": threshold,
        "accuracy": round(float(accuracy_score(labels, predicted)), 4),
        "balanced_accuracy": round(float(balanced_accuracy_score(labels, predicted)), 4),
        "precision": round(float(precision_score(labels, predicted, zero_division=0)), 4),
        "recall": round(float(recall_score(labels, predicted, zero_division=0)), 4),
        "f1": round(float(f1_score(labels, predicted, zero_division=0)), 4),
        "average_precision": (
            round(float(average_precision_score(labels, probabilities)), 4)
            if len(set(labels)) == 2
            else None
        ),
        "brier_score": round(float(brier_score_loss(labels, probabilities)), 4),
        "expected_calibration_error": round(
            _calibration_error(labels, probabilities), 4
        ),
        "confusion_matrix": confusion_matrix(labels, predicted, labels=[0, 1]).tolist(),
        "classification_report": report,
    }


def _fold_metric_summary(folds: list[dict[str, Any]], variant: str) -> dict[str, Any]:
    metric_names = (
        "accuracy", "balanced_accuracy", "precision", "recall", "f1",
        "average_precision", "brier_score", "expected_calibration_error",
    )
    result: dict[str, Any] = {}
    for metric in metric_names:
        values = [
            fold[variant][metric]
            for fold in folds
            if fold[variant][metric] is not None
        ]
        result[metric] = {
            "mean": round(float(np.mean(values)), 4) if values else None,
            "std": round(float(np.std(values)), 4) if values else None,
        }
    return result


def _nested_grouped_evaluation(
    records: list[SignalRecord],
    feature_groups: frozenset[str] | set[str],
    *,
    n_splits: int,
    random_state: int,
) -> dict[str, Any]:
    outer_splits = _fold_count(records, n_splits)
    labels = np.asarray([record.label for record in records], dtype=int)
    groups = _group_ids(records)
    texts = [record.text(feature_groups) for record in records]
    outer = StratifiedGroupKFold(
        n_splits=outer_splits, shuffle=True, random_state=random_state
    )
    raw_oof = np.full(len(records), np.nan, dtype=float)
    calibrated_oof = np.full(len(records), np.nan, dtype=float)
    fold_reports: list[dict[str, Any]] = []

    for fold, (train_indices, test_indices) in enumerate(outer.split(texts, labels, groups)):
        train_records = [records[int(i)] for i in train_indices]
        train_labels = labels[train_indices]
        test_labels = labels[test_indices]
        inner_splits = min(3, _fold_count(train_records, 3))
        inner_raw, _, _ = _oof_raw_probabilities(
            train_records,
            feature_groups,
            n_splits=inner_splits,
            random_state=random_state + fold + 1,
        )
        calibrator = _fit_sigmoid_calibrator(inner_raw, train_labels)

        estimator = _make_pipeline()
        estimator.fit([texts[int(i)] for i in train_indices], train_labels)
        test_raw = estimator.predict_proba([texts[int(i)] for i in test_indices])[:, 1]
        test_calibrated = _calibrated_probabilities(calibrator, test_raw)
        raw_oof[test_indices] = test_raw
        calibrated_oof[test_indices] = test_calibrated

        train_groups = set(groups[train_indices])
        test_groups = set(groups[test_indices])
        fold_reports.append(
            {
                "fold": fold + 1,
                "train_examples": int(len(train_indices)),
                "test_examples": int(len(test_indices)),
                "train_group_count": len(train_groups),
                "test_group_count": len(test_groups),
                "group_overlap_count": len(train_groups & test_groups),
                "raw": _metrics(test_labels, test_raw, 0.5),
                "calibrated": _metrics(test_labels, test_calibrated, 0.5),
            }
        )

    if np.isnan(raw_oof).any() or np.isnan(calibrated_oof).any():
        raise RuntimeError("Nested CV не сформировал полный набор OOF-предсказаний")
    return {
        "outer_folds": outer_splits,
        "raw": {
            "overall": _metrics(labels, raw_oof, 0.5),
            "fold_mean_std": _fold_metric_summary(fold_reports, "raw"),
        },
        "calibrated": {
            "overall": _metrics(labels, calibrated_oof, 0.5),
            "fold_mean_std": _fold_metric_summary(fold_reports, "calibrated"),
        },
        "folds": fold_reports,
    }


def audit_dataset(records: list[SignalRecord]) -> dict[str, Any]:
    """Отчёт о дублях, происхождении и признаках, способных выдавать метку."""
    by_label: dict[str, Any] = {}
    for label in (0, 1):
        subset = [record for record in records if record.label == label]
        feature_modes: dict[str, Any] = {}
        feature_values = {
            "domain": [record.domain or "(пусто)" for record in subset],
            "stage": [stage_bucket(record.stage) for record in subset],
            "trend": [trend_bucket(record.mention_trend) for record in subset],
            "source_count": [source_bucket(record.sources) for record in subset],
        }
        for name, values in feature_values.items():
            counts = Counter(values)
            value, count = counts.most_common(1)[0] if counts else ("(пусто)", 0)
            feature_modes[name] = {
                "most_common": value,
                "count": count,
                "share": round(count / len(subset), 4) if subset else 0.0,
                "distribution": dict(sorted(counts.items())),
            }
        by_label[str(label)] = {
            "examples": len(subset),
            "label_sources": dict(sorted(Counter(r.label_source for r in subset).items())),
            "feature_modes": feature_modes,
        }

    name_counts = Counter(normalize_technology_name(record.technology) for record in records)
    token_counts: dict[str, Counter[str]] = {"0": Counter(), "1": Counter()}
    for record in records:
        tokens = re.findall(
            r"[\w]+",
            record.text(ALL_FEATURE_GROUPS).casefold(),
            flags=re.UNICODE,
        )
        token_counts[str(record.label)].update(set(tokens))
    return {
        "class_distribution": by_label,
        "duplicate_technology_names": {
            name: count for name, count in sorted(name_counts.items()) if count > 1
        },
        "top_tokens_by_class": {
            label: [{"token": token, "records": count} for token, count in counts.most_common(20)]
            for label, counts in token_counts.items()
        },
    }
